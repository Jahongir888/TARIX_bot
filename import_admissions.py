"""Exceldagi o'tish ballarini TargetInstitution va TargetDirection ga yozadi."""
import asyncio
import re
from pathlib import Path

import openpyxl
from sqlalchemy import delete

from database.db import async_session, init_db
from database.models import TargetDirection, TargetInstitution

ROOT = Path(__file__).resolve().parent
EXCEL_CANDIDATES = (
    ROOT / "otish_ballari_fanlar_bilan_5718.xlsx",
    ROOT / "barcha_otish_ballari_5718.xlsx",
)

_APOSTROPHES = ("ʻ", "ʼ", "‘", "’", "`", "´", "ʹ", "′", "‛", "′")


def _norm(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    for ch in _APOSTROPHES:
        text = text.replace(ch, "'")
    text = text.replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", text).strip()


def _to_float(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _to_cipher(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = _norm(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _short_name(name: str) -> str:
    skip = {"va", "hamda", "the", "of", "and"}
    words = re.findall(r"[0-9A-Za-z']+", name)
    initials = "".join(w[0].upper() for w in words if w.lower() not in skip and len(w) > 1)
    return (initials[:10] or name[:12] or "OTM")


def _inst_type(name: str) -> str:
    low = name.lower()
    if "litsey" in low:
        return "litsey_harbiy" if "harbiy" in low else "litsey"
    military = (
        "harbiy",
        "jamoat xavfsizligi",
        "huquqni muhofaza",
        "favqulodda vaziyatlar",
        "milliy gvardiya",
        "ichki ishlar",
        "bojxona",
        "prokuratura",
    )
    if any(key in low for key in military):
        return "otm_harbiy"
    return "otm"


def _find_excel() -> Path:
    for path in EXCEL_CANDIDATES:
        if path.is_file():
            return path
    raise FileNotFoundError(
        "Excel topilmadi. Loyiha ildizida otish_ballari_fanlar_bilan_5718.xlsx "
        "yoki barcha_otish_ballari_5718.xlsx bo'lishi kerak."
    )


def _read_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header = [(_norm(c) if c is not None else "") for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        expected = {
            "region": ("Viloyat/shahar", "region", "viloyat"),
            "university": ("university", "OTM", "otm"),
            "direction": ("direction", "yo'nalish"),
            "edu_form": ("education_form", "ta'lim shakli"),
            "edu_lang": ("education_language", "ta'lim tili"),
            "cipher": ("shifr", "shifri"),
            "grant": ("grant_score", "grant"),
            "contract": ("contract_score", "contract"),
            "main": ("1-fan (Asosiy)", "1-fan", "subject_main"),
            "second": ("2-fan", "subject_second"),
            "mandatory": ("Majburiy fanlar", "mandatory"),
        }
        index = {}
        for key, aliases in expected.items():
            for i, col in enumerate(header):
                if col in aliases or col.lower() in {a.lower() for a in aliases}:
                    index[key] = i
                    break
        missing = [k for k in expected if k not in index]
        if missing:
            raise ValueError(f"Excelda ustunlar yetishmayapti: {', '.join(missing)}. Topilgan: {header}")

        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            university = _norm(raw[index["university"]])
            direction = _norm(raw[index["direction"]])
            if not university or not direction:
                continue
            main = _norm(raw[index["main"]]) or "Noma'lum"
            second = _norm(raw[index["second"]])
            mandatory = _norm(raw[index["mandatory"]])
            block = f"{main} - {second} - MB" if second else f"{main} - MB"
            rows.append({
                "region": _norm(raw[index["region"]]) or "Noma'lum",
                "university": university,
                "direction": direction,
                "edu_form": _norm(raw[index["edu_form"]]),
                "edu_lang": _norm(raw[index["edu_lang"]]) or "O'zbekcha",
                "cipher": _to_cipher(raw[index["cipher"]]),
                "grant": _to_float(raw[index["grant"]]),
                "contract": _to_float(raw[index["contract"]]),
                "main": main,
                "second": second,
                "mandatory": mandatory,
                "block": block,
            })
        return rows
    finally:
        wb.close()


async def import_admissions():
    await init_db()
    path = _find_excel()
    rows = _read_rows(path)
    print(f"Excel: {path.name} — {len(rows)} ta qator o'qildi.")

    async with async_session() as session:
        await session.execute(delete(TargetDirection))
        await session.execute(delete(TargetInstitution))
        await session.commit()

        institutions = {}
        type_counts = {"otm": 0, "otm_harbiy": 0, "litsey": 0, "litsey_harbiy": 0}

        for row in rows:
            inst_type = _inst_type(row["university"])
            key = (row["university"], row["region"], inst_type)
            if key not in institutions:
                inst = TargetInstitution(
                    name=row["university"],
                    short_name=_short_name(row["university"]),
                    inst_type=inst_type,
                    region=row["region"],
                    logo_url=None,
                    is_active=True,
                )
                session.add(inst)
                institutions[key] = inst
            type_counts[inst_type] = type_counts.get(inst_type, 0) + 1

        await session.flush()

        directions = []
        for row in rows:
            inst_type = _inst_type(row["university"])
            inst = institutions[(row["university"], row["region"], inst_type)]
            directions.append(TargetDirection(
                institution_id=inst.id,
                title=row["direction"],
                cipher=row["cipher"],
                edu_form=row["edu_form"],
                edu_lang=row["edu_lang"],
                education_lang=row["edu_lang"],
                subject_main=row["main"],
                subject_second=row["second"],
                subject_secondary=row["second"],
                mandatory_subjects=row["mandatory"],
                subject_block=row["block"],
                grant_score=row["grant"],
                contract_score=row["contract"],
                grant_ball=row["grant"],
                contract_ball=row["contract"],
                quota=0,
                mock_quiz_id=1,
            ))

        session.add_all(directions)
        await session.commit()

    print(f"Muassasalar: {len(institutions)}")
    print(f"Yo'nalishlar: {len(directions)}")
    print("Toifa bo'yicha qatorlar:", ", ".join(f"{k}={v}" for k, v in type_counts.items()))
    print("Import tugadi.")


if __name__ == "__main__":
    asyncio.run(import_admissions())
