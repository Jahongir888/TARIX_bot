"""harbiy_muassasalar.xlsx dagi harbiy litsey va OTMlarni bazaga yuklaydi.

Faqat inst_type = litsey_harbiy | otm_harbiy yozuvlarini yangilaydi
(fuqarolik OTM/litsey yozuvlariga tegmaydi).

Maydonlar:
  TargetInstitution — inst_type, region, logo_url, gallery_images, review_video_url,
                      overview_text, eligibility_criteria, career_opportunities, admission_info
  TargetDirection — direction_title, subject_main/second/third, grant_score, mock_test_url,
                    catalog_courses_url, exam1–exam5, dates_timeline_json, practical_exams_json
"""
from __future__ import annotations

import asyncio
import json
import re
from datetime import datetime, date
from pathlib import Path

import openpyxl
from sqlalchemy import delete, select

from database.db import async_session, init_db
from database.models import TargetDirection, TargetInstitution

ROOT = Path(__file__).resolve().parent
EXCEL_PATH = ROOT / "harbiy_muassasalar.xlsx"
MILITARY_TYPES = ("litsey_harbiy", "otm_harbiy")

_APOSTROPHES = ("ʻ", "ʼ", "‘", "’", "`", "´", "ʹ", "′", "‛")
_URL_RE = re.compile(r"^https?://", re.I)

# Excel ustuni -> ichki kalit (harbiy_muassasalar.xlsx sarlavhalari)
COLUMN_ALIASES = {
    "inst_type": ("inst_type",),
    "name": ("inst_full_name", "name", "nomi", "muassasa"),
    "short_name": ("inst_short_name", "short_name", "qisqa_nomi", "qisqa nomi"),
    "title": ("direction_title", "title", "yo'nalish", "yonalish"),
    "logo_url": ("logo_url", "logo", "logosi"),
    "region": ("region", "viloyat", "viloyat/shahar"),
    "gallery": ("gallery_images", "gallery", "galereya", "galereyasi"),
    "video_review_url": ("review_video_url", "video_review_url", "videosharh"),
    "overview_text": ("overview_text", "umumiy", "tavsif"),
    "eligibility_criteria": ("eligibility_criteria", "qabul_talablari", "talablar"),
    "opportunities": ("career_opportunities", "opportunities", "imkoniyatlar", "imkoniyatlari"),
    "admission_info": ("admission_info",),
    "admission_date": ("admission_date",),
    "subject_main": ("subject_main", "1-fan"),
    "subject_second": ("subject_second", "2-fan"),
    "subject_third": ("subject_third", "subject3_exam", "3-fan", "3_fan"),
    "grant_score": ("grant_score", "grant"),
    "subjects_table_image": ("subjects_table_image",),
    "mock_test_url": ("mock_test_url",),
    "catalog_courses_url": ("catalog_courses_url",),
    "bundle_package_url": ("bundle_package_url", "package_url", "paket_url"),
    "practical_mock_id": ("practical_mock_id",),
    "practical_course_url": ("practical_course_url",),
    "test_exam_date": ("test_exam_date",),
}

for n in range(1, 6):
    COLUMN_ALIASES[f"exam{n}_date"] = (f"exam{n}_date",)
    COLUMN_ALIASES[f"exam{n}_name_criteria"] = (f"exam{n}_name_criteria",)
    COLUMN_ALIASES[f"exam{n}_name"] = (f"exam{n}_name",)
    COLUMN_ALIASES[f"exam{n}_criteria"] = (f"exam{n}_criteria",)
    COLUMN_ALIASES[f"exam{n}_mock_url"] = (f"exam{n}_mock_url", f"exam{n}_mock")
    COLUMN_ALIASES[f"exam{n}_course_url"] = (f"exam{n}_course_url", f"exam{n}_course")


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for ch in _APOSTROPHES:
        text = text.replace(ch, "'")
    text = text.replace("“", '"').replace("”", '"')
    return text.strip()


def _norm_header(value) -> str:
    return _norm(value).lower().replace(" ", "_")


def _to_float(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", ".").strip())
    except (TypeError, ValueError):
        return 0.0


def _to_int(value, default: int = 1) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def _short_name(name: str, explicit: str = "") -> str:
    if explicit:
        return explicit[:64]
    skip = {"va", "hamda", "the", "of", "and", "huzuridagi", "nomidagi"}
    words = re.findall(r"[0-9A-Za-z']+", name)
    initials = "".join(w[0].upper() for w in words if w.lower() not in skip and len(w) > 1)
    return (initials[:10] or name[:12] or "HARBIY")


def _normalize_inst_type(raw: str, name: str = "") -> str | None:
    text = _norm(raw).lower().replace(" ", "_").replace("-", "_")
    aliases = {
        "litsey_harbiy": "litsey_harbiy",
        "harbiy_litsey": "litsey_harbiy",
        "otm_harbiy": "otm_harbiy",
        "harbiy_otm": "otm_harbiy",
    }
    if text in aliases:
        return aliases[text]
    low = f"{text} {name}".lower()
    if "litsey" in low:
        return "litsey_harbiy"
    if "harbiy" in low or "akademiya" in low:
        return "otm_harbiy"
    return None


def _split_name_criteria(raw: str, default_name: str | None = None) -> tuple[str | None, str | None]:
    """examN_name_criteria: nom+mezon, yoki faqat mezon rasmi URL."""
    text = _norm(raw)
    if not text:
        return None, None
    # Faqat URL bo'lsa — mezon rasmi / hujjat
    if _URL_RE.match(text) and "\n" not in text and "|" not in text:
        return default_name, text
    parts = [p.strip() for p in re.split(r"[\n|;]+", text) if p.strip()]
    if len(parts) == 1:
        for sep in (" — ", " - ", ": ", " – "):
            if sep in parts[0]:
                name, _, criteria = parts[0].partition(sep)
                return name.strip() or default_name, criteria.strip() or None
        return parts[0], None
    return parts[0], "\n".join(parts[1:])


def _build_timeline_and_exams(row: dict) -> tuple[str, str]:
    """dates_timeline va practical_exams ni JSON string sifatida yig'adi."""
    timeline = []
    if row.get("admission_date"):
        timeline.append({
            "key": "admission",
            "label": "Hujjat topshirish",
            "date": row["admission_date"],
        })
    exams = []
    for n in range(1, 6):
        exam_date = row.get(f"exam{n}_date")
        exam_name = row.get(f"exam{n}_name")
        exam_criteria = row.get(f"exam{n}_criteria")
        mock_url = row.get(f"exam{n}_mock_url")
        course_url = row.get(f"exam{n}_course_url")
        if not any([exam_date, exam_name, exam_criteria, mock_url, course_url]):
            continue
        label = exam_name or f"{n}-amaliy imtihon"
        if exam_date:
            timeline.append({"key": f"exam{n}", "label": label, "date": exam_date})
        exams.append({
            "index": n,
            "name": label,
            "date": exam_date or "",
            "criteria": exam_criteria or "",
            "practical_mock_url": mock_url,
            "practical_course_url": course_url,
            "practical_mock_id": row.get("practical_mock_id"),
        })
    if row.get("test_exam_date"):
        timeline.append({
            "key": "test_exam",
            "label": "Test sinovi",
            "date": row["test_exam_date"],
        })
    return json.dumps(timeline, ensure_ascii=False), json.dumps(exams, ensure_ascii=False)


def _build_admission_requirements(eligibility: str, admission_info: str, admission_date: str) -> str | None:
    chunks = []
    if eligibility:
        chunks.append(eligibility)
    if admission_info:
        chunks.append("Qabul tartibi:\n" + admission_info)
    if admission_date:
        chunks.append(f"Qabul sanasi: {admission_date}")
    return "\n\n".join(chunks) if chunks else None


def _cell(raw, index: dict, key: str) -> str:
    if key not in index:
        return ""
    idx = index[key]
    if idx >= len(raw):
        return ""
    return _norm(raw[idx])


def _build_header_index(header: list[str]) -> dict:
    normalized = [_norm_header(h) for h in header]
    index = {}
    for key, aliases in COLUMN_ALIASES.items():
        alias_set = {_norm_header(a) for a in aliases}
        for i, col in enumerate(normalized):
            if col in alias_set:
                index[key] = i
                break
    return index


def _read_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        header = [(_norm(c) if c is not None else "") for c in header_row]
        index = _build_header_index(header)
        if "name" not in index:
            raise ValueError(
                f"Excelda 'inst_full_name' / 'name' ustuni topilmadi. Topilgan: {header}"
            )

        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            name = _cell(raw, index, "name")
            if not name:
                continue

            inst_type = _normalize_inst_type(_cell(raw, index, "inst_type"), name)
            if inst_type not in MILITARY_TYPES:
                print(f"O'tkazib yuborildi (inst_type emas): {name!r} -> {_cell(raw, index, 'inst_type')!r}")
                continue

            main = _cell(raw, index, "subject_main") or "Tarix"
            second = _cell(raw, index, "subject_second")
            third = _cell(raw, index, "subject_third")
            if second and third:
                block = f"{main} - {second} - {third}"
            elif second:
                block = f"{main} - {second} - MB"
            else:
                block = f"{main} - MB"

            title = _cell(raw, index, "title") or (
                "Harbiy tayyorgarlik" if inst_type == "otm_harbiy" else "Ijtimoiy-gumanitar fanlar"
            )

            practical_course = _cell(raw, index, "practical_course_url") or None
            practical_mock_raw = _cell(raw, index, "practical_mock_id") or None
            mock_quiz_id = 1
            shared_mock_url = None
            if practical_mock_raw:
                if practical_mock_raw.startswith("http"):
                    shared_mock_url = practical_mock_raw
                else:
                    mock_quiz_id = _to_int(practical_mock_raw, 1)

            eligibility = _cell(raw, index, "eligibility_criteria") or None
            admission_info = _cell(raw, index, "admission_info") or None
            admission_date = _cell(raw, index, "admission_date") or None
            catalog_url = _cell(raw, index, "catalog_courses_url") or None
            bundle_url = _cell(raw, index, "bundle_package_url") or catalog_url
            overview = _cell(raw, index, "overview_text") or None
            opportunities = _cell(raw, index, "opportunities") or None

            row = {
                "name": name,
                "short_name": _short_name(name, _cell(raw, index, "short_name")),
                "inst_type": inst_type,
                "region": _cell(raw, index, "region") or "Noma'lum",
                "logo_url": _cell(raw, index, "logo_url") or None,
                "gallery": _cell(raw, index, "gallery") or None,
                "video_review_url": _cell(raw, index, "video_review_url") or None,
                "eligibility_criteria": eligibility,
                "admission_info": admission_info,
                "admission_date": admission_date,
                "admission_requirements": _build_admission_requirements(
                    eligibility or "",
                    admission_info or "",
                    admission_date or "",
                ),
                "opportunities": opportunities,
                "overview_text": overview,
                "title": title,
                "subject_main": main,
                "subject_second": second or None,
                "subject_third": third or None,
                "subject_block": block,
                "grant_score": _to_float(raw[index["grant_score"]]) if "grant_score" in index else 0.0,
                "mock_test_url": _cell(raw, index, "mock_test_url") or None,
                "catalog_courses_url": catalog_url,
                "bundle_package_url": bundle_url,
                "subjects_table_image": _cell(raw, index, "subjects_table_image") or None,
                "test_exam_date": _cell(raw, index, "test_exam_date") or None,
                "practical_mock_id": practical_mock_raw,
                "practical_course_url": practical_course,
                "mock_quiz_id": mock_quiz_id,
            }

            for n in range(1, 6):
                default_name = f"{n}-amaliy imtihon"
                exam_date = _cell(raw, index, f"exam{n}_date") or None
                explicit_name = _cell(raw, index, f"exam{n}_name") or None
                explicit_criteria = _cell(raw, index, f"exam{n}_criteria") or None
                split_name, split_criteria = _split_name_criteria(
                    _cell(raw, index, f"exam{n}_name_criteria"),
                    default_name=default_name,
                )
                exam_name = explicit_name or split_name
                exam_criteria = explicit_criteria or split_criteria
                per_mock = _cell(raw, index, f"exam{n}_mock_url") or None
                per_course = _cell(raw, index, f"exam{n}_course_url") or None
                has_exam = bool(exam_date or exam_name or exam_criteria or per_mock or per_course)
                row[f"exam{n}_date"] = exam_date
                row[f"exam{n}_name"] = exam_name if has_exam else None
                row[f"exam{n}_criteria"] = exam_criteria
                row[f"exam{n}_mock_url"] = per_mock or (shared_mock_url if has_exam else None)
                row[f"exam{n}_course_url"] = per_course or (practical_course if has_exam else None)

            timeline_json, exams_json = _build_timeline_and_exams(row)
            row["dates_timeline_json"] = timeline_json
            row["practical_exams_json"] = exams_json
            rows.append(row)
        return rows
    finally:
        wb.close()


async def import_military():
    await init_db()
    if not EXCEL_PATH.is_file():
        raise FileNotFoundError(
            f"Excel topilmadi: {EXCEL_PATH}\n"
            "Loyiha ildiziga harbiy_muassasalar.xlsx qo'ying."
        )

    rows = _read_rows(EXCEL_PATH)
    print(f"Excel: {EXCEL_PATH.name} — {len(rows)} ta harbiy qator o'qildi.")
    if not rows:
        print("Yoziladigan qator yo'q. Import to'xtatildi.")
        return

    async with async_session() as session:
        existing_ids = (
            await session.execute(
                select(TargetInstitution.id).where(TargetInstitution.inst_type.in_(MILITARY_TYPES))
            )
        ).scalars().all()
        if existing_ids:
            await session.execute(
                delete(TargetDirection).where(TargetDirection.institution_id.in_(existing_ids))
            )
            await session.execute(
                delete(TargetInstitution).where(TargetInstitution.id.in_(existing_ids))
            )
            await session.commit()
            print(f"Eski harbiy yozuvlar o'chirildi: {len(existing_ids)} ta muassasa.")

        type_counts = {"litsey_harbiy": 0, "otm_harbiy": 0}
        institutions = {}

        for row in rows:
            key = (row["name"], row["region"], row["inst_type"])
            if key not in institutions:
                inst = TargetInstitution(
                    name=row["name"],
                    short_name=row["short_name"],
                    inst_type=row["inst_type"],
                    region=row["region"],
                    logo_url=row["logo_url"],
                    gallery=row["gallery"],
                    video_review_url=row["video_review_url"],
                    overview_text=row.get("overview_text"),
                    admission_requirements=row["admission_requirements"],
                    eligibility_criteria=row["eligibility_criteria"],
                    admission_info=row["admission_info"],
                    admission_date=row["admission_date"],
                    opportunities=row["opportunities"],
                    is_active=True,
                )
                session.add(inst)
                institutions[key] = inst
            type_counts[row["inst_type"]] = type_counts.get(row["inst_type"], 0) + 1

        await session.flush()

        directions = []
        for row in rows:
            key = (row["name"], row["region"], row["inst_type"])
            inst = institutions[key]
            exam_kwargs = {
                f"exam{n}_{suffix}": row.get(f"exam{n}_{suffix}")
                for n in range(1, 6)
                for suffix in ("date", "name", "criteria", "mock_url", "course_url")
            }
            directions.append(TargetDirection(
                institution_id=inst.id,
                title=row["title"],
                subject_main=row["subject_main"],
                subject_second=row["subject_second"],
                subject_third=row["subject_third"],
                subject_secondary=row["subject_second"],
                subject_block=row["subject_block"],
                education_lang="O'zbek",
                edu_lang="O'zbek",
                grant_score=row["grant_score"],
                grant_ball=row["grant_score"],
                contract_score=0.0,
                contract_ball=0.0,
                mock_quiz_id=row["mock_quiz_id"],
                mock_test_url=row["mock_test_url"],
                catalog_courses_url=row["catalog_courses_url"],
                bundle_package_url=row["bundle_package_url"],
                subjects_table_image=row["subjects_table_image"],
                test_exam_date=row["test_exam_date"],
                practical_mock_id=row["practical_mock_id"],
                practical_course_url=row["practical_course_url"],
                dates_timeline_json=row.get("dates_timeline_json"),
                practical_exams_json=row.get("practical_exams_json"),
                **exam_kwargs,
            ))

        session.add_all(directions)
        await session.commit()

    print(f"Muassasalar: {len(institutions)}")
    print(f"Yo'nalishlar: {len(directions)}")
    print("Toifa:", ", ".join(f"{k}={v}" for k, v in type_counts.items()))
    print("Import tugadi.")


if __name__ == "__main__":
    asyncio.run(import_military())
