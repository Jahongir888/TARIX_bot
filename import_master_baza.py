"""master_baza.xlsx (1-sahifa) ni SQLite bazaga xavfsiz import qiladi.

Excel fayl HECH QANDAY o'zgartirilmaydi — mavjud ustunlar bilan o'qiladi.
openpyxl ishlatiladi (serverda pandas shart emas).
"""
from __future__ import annotations

import asyncio
import math
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

from database.db import async_session, init_db
from database.models import MasterFactQuestion

ROOT = Path(__file__).resolve().parent
EXCEL_PATH = ROOT / "master_baza.xlsx"

EXPECTED_COLUMNS = [
    "Fakt_ID", "Test_ID", "Asl_Gap", "Fakt_Kontekst", "Savol_Matni",
    "Test_Matni", "Test_Obyekt", "Etalon_Javob", "Variant_B", "Variant_C",
    "Variant_D", "Kalit_Soz", "Klass", "Yonalish", "Tarkibiy_Faktlar",
    "Tematik_Teg", "Mavzu_Reja",
]

# Fakt_ID / Test_ID / Mavzu_Reja: 6_01_01 | T_6_01_01 | R_6_01_01
_ID_RE = re.compile(
    r"^(?:[TR]_)?(?P<grade>\d+)_(?P<topic>\d+)(?:_(?P<sub>\d+))?$",
    re.IGNORECASE,
)


def _cell(value) -> str | None:
    """Excel katakchasini tozalangan matnga aylantiradi; bo'sh bo'lsa None."""
    if value is None:
        return None
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return None
    return text


def _parse_grade_topic(raw_id: str | None) -> tuple[int | None, int | None]:
    """'6_01_01' / 'T_6_01_01' / 'R_6_01_01' dan (grade, topic_num) ajratadi."""
    if not raw_id:
        return None, None
    text = raw_id.strip()
    m = _ID_RE.match(text)
    if not m:
        return None, None
    try:
        return int(m.group("grade")), int(m.group("topic"))
    except (TypeError, ValueError):
        return None, None


def _load_rows() -> list[dict]:
    if not EXCEL_PATH.is_file():
        raise FileNotFoundError(f"Excel topilmadi: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        raise ValueError("Excel 1-sahifasi bo'sh")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    col_index = {name: headers.index(name) for name in EXPECTED_COLUMNS if name in headers}
    missing = [c for c in EXPECTED_COLUMNS if c not in col_index]
    if missing:
        wb.close()
        raise ValueError(f"Excel ustunlari yetishmayapti: {missing}")

    def get(row_vals, name: str):
        idx = col_index[name]
        return row_vals[idx] if idx < len(row_vals) else None

    rows: list[dict] = []
    last_fakt_id: str | None = None
    last_asl_gap: str | None = None
    last_grade: int | None = None
    last_topic: int | None = None
    skipped = 0
    total = 0

    for row_vals in rows_iter:
        total += 1
        if row_vals is None:
            skipped += 1
            continue
        row_vals = list(row_vals)

        fakt_id = _cell(get(row_vals, "Fakt_ID"))
        test_id = _cell(get(row_vals, "Test_ID"))
        asl_gap = _cell(get(row_vals, "Asl_Gap"))
        savol = _cell(get(row_vals, "Savol_Matni"))
        test_matni = _cell(get(row_vals, "Test_Matni"))
        etalon = _cell(get(row_vals, "Etalon_Javob"))
        mavzu_reja = _cell(get(row_vals, "Mavzu_Reja"))

        # Ikkala ID ham bo'sh — nazariy qo'shimcha yoki bo'sh satr
        if not fakt_id and not test_id:
            skipped += 1
            continue

        # Fakt_ID bo'sh, Test_ID bor — oldingi to'liq faktga bog'laymiz
        if not fakt_id:
            fakt_id = last_fakt_id
            if not fakt_id and test_id:
                fakt_id = test_id[2:] if test_id.upper().startswith("T_") else test_id
            if not asl_gap:
                asl_gap = last_asl_gap

        grade, topic_num = _parse_grade_topic(fakt_id)
        if grade is None or topic_num is None:
            grade, topic_num = _parse_grade_topic(test_id)
        if grade is None or topic_num is None:
            grade, topic_num = _parse_grade_topic(mavzu_reja)
        if grade is None:
            grade = last_grade
        if topic_num is None:
            topic_num = last_topic

        if not asl_gap:
            asl_gap = last_asl_gap

        question_text = savol or test_matni
        if not fakt_id or grade is None or topic_num is None or not asl_gap or not question_text or not etalon:
            skipped += 1
            continue

        last_fakt_id = fakt_id
        last_asl_gap = asl_gap
        last_grade = grade
        last_topic = topic_num

        rows.append({
            "fakt_id": fakt_id,
            "test_id": test_id,
            "grade": grade,
            "topic_num": topic_num,
            "asl_gap": asl_gap,
            "fakt_kontekst": _cell(get(row_vals, "Fakt_Kontekst")),
            "savol_matni": question_text,
            "test_matni": test_matni,
            "test_obyekt": _cell(get(row_vals, "Test_Obyekt")),
            "etalon_javob": etalon,
            "variant_b": _cell(get(row_vals, "Variant_B")),
            "variant_c": _cell(get(row_vals, "Variant_C")),
            "variant_d": _cell(get(row_vals, "Variant_D")),
            "kalit_soz": _cell(get(row_vals, "Kalit_Soz")),
            "klass": _cell(get(row_vals, "Klass")),
            "yonalish": _cell(get(row_vals, "Yonalish")),
            "tarkibiy_faktlar": _cell(get(row_vals, "Tarkibiy_Faktlar")),
            "tematik_teg": _cell(get(row_vals, "Tematik_Teg")),
            "mavzu_reja": mavzu_reja,
        })

    wb.close()
    print(f"[OK] Excel: {total} satr o'qildi, {len(rows)} ta yozuv tayyor, {skipped} ta o'tkazib yuborildi.")
    return rows


async def import_master_baza(*, replace: bool = True) -> None:
    await init_db()
    records = _load_rows()
    if not records:
        print("[WARN] Import qilinadigan yozuv topilmadi.")
        return

    async with async_session() as session:
        async with session.begin():
            if replace:
                existing = (
                    await session.execute(select(MasterFactQuestion.id).limit(1))
                ).scalar_one_or_none()
                if existing is not None:
                    await session.execute(delete(MasterFactQuestion))
                    print("[OK] Eski master_fact_questions yozuvlari o'chirildi.")

            session.add_all([MasterFactQuestion(**item) for item in records])

        facts = len({r["fakt_id"] for r in records})
        questions = len(records)
        grades = sorted({r["grade"] for r in records})
        topics = sorted({(r["grade"], r["topic_num"]) for r in records})
        print(f"[OK] Yuklandi: {facts} ta fakt, {questions} ta savol.")
        print(f"     Sinflar: {grades}")
        print(f"     Mavzular: {len(topics)} ta (grade, topic_num) juftlik.")


if __name__ == "__main__":
    asyncio.run(import_master_baza())
