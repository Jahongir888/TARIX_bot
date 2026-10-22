"""Admin Excel orqali yangiliklar importi (/add_news)."""
import os
from typing import Any, Optional

from aiogram import Router, F, Bot
from aiogram.dispatcher.event.bases import SkipHandler
from aiogram.filters import Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message

from data.config import ADMINS
from database.db import async_session
from database.models import News

admin_news_router = Router()

NEWS_ADMINS = list(dict.fromkeys([*(ADMINS or []), 1206236612]))


class ImportNewsState(StatesGroup):
    waiting_file = State()


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _cell_str(value: Any, default: str = "") -> str:
    if _is_blank(value):
        return default
    return str(value).strip()


def _cell_optional_str(value: Any) -> Optional[str]:
    if _is_blank(value):
        return None
    return str(value).strip()


def _cell_optional_int(value: Any) -> Optional[int]:
    if _is_blank(value):
        return None
    return int(value)


def _is_excel_name(name: str) -> bool:
    lower = (name or "").lower()
    return lower.endswith(".xlsx") or lower.endswith(".xls")


def _is_quiz_workbook(path: str) -> bool:
    """Test importi bilan chalkashmaslik: CONFIG+QUESTIONS bo'lsa — yangilik emas."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet_names = {str(name).strip().upper() for name in wb.sheetnames}
        wb.close()
        return "CONFIG" in sheet_names and "QUESTIONS" in sheet_names
    except Exception:
        return False


def _read_news_rows(file_path: str) -> list[dict]:
    """Excelning birinchi varagidan header+qatorlarni o'qiydi (pandas o'rniga)."""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_raw = next(rows_iter, None)
        if not header_raw:
            return []
        headers = [
            str(cell).strip().lower() if cell is not None else ""
            for cell in header_raw
        ]
        records: list[dict] = []
        for raw in rows_iter:
            if raw is None or all(_is_blank(cell) for cell in raw):
                continue
            item: dict = {}
            for index, key in enumerate(headers):
                if not key:
                    continue
                item[key] = raw[index] if index < len(raw) else None
            records.append(item)
        return records
    finally:
        wb.close()


async def _download_document(message: Message, bot: Bot) -> str:
    os.makedirs("downloads", exist_ok=True)
    filename = message.document.file_name or "news.xlsx"
    safe_name = os.path.basename(filename)
    file_path = os.path.join("downloads", safe_name)
    file = await bot.get_file(message.document.file_id)
    await bot.download_file(file.file_path, file_path)
    return file_path


async def _import_news_workbook(file_path: str, bot: Bot) -> int:
    rows = _read_news_rows(file_path)
    count = 0
    async with async_session() as session:
        try:
            for row in rows:
                title = _cell_str(row.get("title"))
                content = _cell_str(row.get("content"))
                if not title or not content:
                    continue
                new_item = News(
                    title=title,
                    hashtags=_cell_str(row.get("hashtags")),
                    category=_cell_str(row.get("category"), "umumiy") or "umumiy",
                    summary=_cell_optional_str(row.get("summary")),
                    content=content,
                    image_url=_cell_str(row.get("image_url")),
                    source_url=_cell_optional_str(row.get("source_url")),
                    source_name=_cell_str(row.get("source_name"), "Manba") or "Manba",
                    discussion_chat_id=_cell_optional_int(row.get("discussion_chat_id")),
                )
                session.add(new_item)
                await session.flush()

                if new_item.discussion_chat_id:
                    try:
                        sent = await bot.send_photo(
                            chat_id=new_item.discussion_chat_id,
                            photo=new_item.image_url,
                            caption=(
                                f"📢 <b>{new_item.title}</b>\n\n"
                                f"{new_item.summary or ''}\n\n{new_item.hashtags}"
                            ),
                            parse_mode="HTML",
                        )
                        new_item.telegram_post_msg_id = sent.message_id
                    except Exception as err:
                        print(f"Guruhga post chiqarishda xatolik: {err}")

                count += 1
            await session.commit()
        except Exception:
            await session.rollback()
            raise
    return count


@admin_news_router.message(Command("add_news"), F.from_user.id.in_(NEWS_ADMINS))
async def cmd_add_news(message: Message, state: FSMContext):
    await state.set_state(ImportNewsState.waiting_file)
    await message.answer(
        "📰 Yangiliklar Excelini yuboring.\n\n"
        "Ustunlar: <b>title</b>, <b>content</b> (majburiy);\n"
        "ixtiyoriy: hashtags, category, summary, image_url, "
        "source_url, source_name, discussion_chat_id.\n\n"
        "Bekor qilish: /start",
        parse_mode="HTML",
    )


@admin_news_router.message(ImportNewsState.waiting_file, F.document, F.from_user.id.in_(NEWS_ADMINS))
async def add_news_waiting_file(message: Message, bot: Bot, state: FSMContext):
    if not _is_excel_name(message.document.file_name or ""):
        await message.answer("❌ .xlsx yoki .xls fayl yuboring.")
        return
    status = await message.answer("📥 Yangiliklar import qilinmoqda...")
    file_path = await _download_document(message, bot)
    try:
        if _is_quiz_workbook(file_path):
            await status.edit_text(
                "❌ Bu test shabloni (CONFIG/QUESTIONS). "
                "Yangilik uchun boshqa Excel yuboring yoki /import_quiz ishlating."
            )
            return
        count = await _import_news_workbook(file_path, bot)
        await state.clear()
        await status.edit_text(
            f"✅ Muvaffaqiyatli <b>{count}</b> ta yangilik bazaga kiritildi!",
            parse_mode="HTML",
        )
    except Exception as e:
        await status.edit_text(f"❌ Xatolik: {e}")
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)


@admin_news_router.message(StateFilter(None), F.document, F.from_user.id.in_(NEWS_ADMINS))
async def import_news_excel(message: Message, bot: Bot):
    """Hujjat yuborilganda (quiz bo'lmasa) avtomatik yangilik importi."""
    if not _is_excel_name(message.document.file_name or ""):
        return

    file_path = await _download_document(message, bot)
    try:
        if _is_quiz_workbook(file_path):
            # Test importi admin_quiz_import da ishlaydi
            raise SkipHandler()

        status_msg = await message.answer("📥 Excel fayl qabul qilindi, yangiliklar import qilinmoqda...")
        count = await _import_news_workbook(file_path, bot)
        await status_msg.edit_text(
            f"✅ Muvaffaqiyatli <b>{count}</b> ta yangilik bazaga kiritildi!",
            parse_mode="HTML",
        )
    except SkipHandler:
        raise
    except Exception as e:
        await message.answer(f"❌ Xatolik yuz berdi: {str(e)}")
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)
