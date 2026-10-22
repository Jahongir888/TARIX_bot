"""Admin Excel orqali namunaviy test va milliy sertifikat importi.

Fayl ikki sahifadan iborat: CONFIG va QUESTIONS.
"""
import html
import math
import os
from datetime import datetime, date

import openpyxl
from aiogram import Router, F, Bot
from aiogram.dispatcher.event.bases import SkipHandler
from aiogram.filters import Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message, InlineKeyboardMarkup, InlineKeyboardButton, WebAppInfo
from sqlalchemy import delete, select

from data.config import ADMINS, WEBAPP_URL
from database.db import async_session
from database.models import ExamQuiz, ExamQuestion

admin_quiz_router = Router()

# admin_news dagi admin bilan bir xil zaxira, .env dagi ADMINS ustun.
QUIZ_ADMINS = list(dict.fromkeys([*(ADMINS or []), 1206236612]))

SUBJECT_ALIASES = {
    "tarix": "tarix",
    "history": "tarix",
    "huquq": "huquq",
    "huquqshunoslik": "huquq",
    "psixologiya": "psixologiya",
    "psix": "psixologiya",
    "ona_tili": "ona_tili",
    "ona-tili": "ona_tili",
    "onatili": "ona_tili",
    "ona tili": "ona_tili",
}

EXAM_TYPE_ALIASES = {
    "ms": "milliy_sertifikat",
    "milliy": "milliy_sertifikat",
    "milliy_sertifikat": "milliy_sertifikat",
    "milliy-sertifikat": "milliy_sertifikat",
    "namunaviy": "namunaviy",
    "namuna": "namunaviy",
    "imtihon": "imtihon",
    "bmba": "bmba_asosiy",
    "bmba_asosiy": "bmba_asosiy",
    "bmba_majburiy": "bmba_majburiy",
    "olimpiada": "olimpiada",
    "attestatsiya": "attestatsiya",
    "akademik_litsey": "akademik_litsey",
    "litsey": "akademik_litsey",
}

CONFIG_KEYS = {
    "command_path", "quiz_title", "title", "visible_at", "starts_at",
    "base_price", "free_for_courses", "share_discount_percent",
    "full_analysis_video", "subject", "exam_type", "duration_minutes",
    "link", "analysis_text",
}
QUESTION_KEYS = {
    "order_num", "question_text", "image_top", "image_bottom",
    "opt_a", "opt_b", "opt_c", "opt_d", "option_a", "option_b", "option_c", "option_d",
    "correct", "correct_option", "beta", "difficulty_beta",
    "ref_a", "ref_b", "ref_c", "ref_d", "solution_text", "video_url", "video_solution_url",
}


class ImportQuizState(StatesGroup):
    waiting_file = State()


def _norm_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _as_str(value):
    if _is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _as_int(value, default=0) -> int:
    if _is_blank(value):
        return default
    return int(float(value))


def _as_float(value, default=2.0) -> float:
    if _is_blank(value):
        return default
    return float(value)


def _parse_dt(value):
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"vaqt formati tushunilmadi: {text}")


def _clean_courses(value) -> str:
    text = _as_str(value)
    if not text:
        return ""
    parts = []
    for part in text.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        try:
            num = float(part)
            if num.is_integer():
                parts.append(str(int(num)))
                continue
        except ValueError:
            pass
        parts.append(part)
    return ",".join(parts)


def _correct_index(value) -> int:
    if _is_blank(value):
        raise ValueError("correct ustuni bo'sh")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = int(value)
        if number == 0:
            return 0
        if number in (1, 2, 3, 4):
            return number - 1
        raise ValueError(f"correct qiymati noto'g'ri: {value}")
    text = str(value).strip().upper()
    if text in ("A", "B", "C", "D"):
        return "ABCD".index(text)
    if text in ("1", "2", "3", "4"):
        return int(text) - 1
    if text == "0":
        return 0
    raise ValueError(f"correct qiymati noto'g'ri: {value}")


def _clamp_beta(value) -> float:
    beta = _as_float(value, 2.0)
    if beta < 1.1:
        return 1.1
    if beta > 3.2:
        return 3.2
    return round(beta, 3)


def is_quiz_workbook(path: str) -> bool:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        names = {str(name).strip().upper() for name in wb.sheetnames}
    finally:
        wb.close()
    return "CONFIG" in names and "QUESTIONS" in names


def _sheet_by_name(wb, name: str):
    for sheet in wb.worksheets:
        if str(sheet.title).strip().upper() == name:
            return sheet
    return None


def _sheet_records(ws) -> list:
    rows = []
    for raw in ws.iter_rows(values_only=True):
        if any(not _is_blank(cell) for cell in raw):
            rows.append(list(raw))
    if not rows:
        return []
    headers = [_norm_header(cell) for cell in rows[0]]
    header_set = {key for key in headers if key}
    if header_set & (CONFIG_KEYS | QUESTION_KEYS):
        records = []
        for raw in rows[1:]:
            item = {}
            for index, key in enumerate(headers):
                if not key:
                    continue
                item[key] = raw[index] if index < len(raw) else None
            if any(not _is_blank(val) for val in item.values()):
                records.append(item)
        return records
    item = {}
    for raw in rows:
        key = _norm_header(raw[0]) if raw else ""
        if not key:
            continue
        item[key] = raw[1] if len(raw) > 1 else None
    return [item] if item else []


def _infer_subject_and_type(command_path: str, subject_raw, exam_type_raw) -> tuple:
    subject = None
    exam_type = None
    explicit_subject = _norm_header(subject_raw)
    if explicit_subject:
        subject = (
            SUBJECT_ALIASES.get(explicit_subject)
            or SUBJECT_ALIASES.get(explicit_subject.replace("_", " "))
            or explicit_subject
        )
    explicit_type = _norm_header(exam_type_raw)
    if explicit_type:
        exam_type = EXAM_TYPE_ALIASES.get(explicit_type, explicit_type)

    parts = [part for part in (command_path or "").strip().strip("/").split("/") if part]
    for part in parts:
        token = _norm_header(part)
        if exam_type is None and token in EXAM_TYPE_ALIASES:
            exam_type = EXAM_TYPE_ALIASES[token]
        if subject is None and token in SUBJECT_ALIASES:
            subject = SUBJECT_ALIASES[token]
    return subject or "tarix", exam_type or "namunaviy"


def _option_sources(row: dict) -> str | None:
    import json
    sources = {}
    for key, column in (("A", "ref_a"), ("B", "ref_b"), ("C", "ref_c"), ("D", "ref_d")):
        text = _as_str(row.get(column))
        if text:
            sources[key] = text
    if not sources:
        return None
    return json.dumps(sources, ensure_ascii=False)


def parse_quiz_workbook(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        config_sheet = _sheet_by_name(wb, "CONFIG")
        question_sheet = _sheet_by_name(wb, "QUESTIONS")
        if config_sheet is None or question_sheet is None:
            raise ValueError("Excelda CONFIG va QUESTIONS sahifalari bo'lishi kerak")
        config_rows = _sheet_records(config_sheet)
        if not config_rows:
            raise ValueError("CONFIG sahifasi bo'sh")
        config = config_rows[0]
        title = _as_str(config.get("quiz_title") or config.get("title"))
        if not title:
            raise ValueError("CONFIG da quiz_title yo'q")
        command_path = _as_str(config.get("command_path"))
        if command_path and not command_path.startswith("/"):
            command_path = "/" + command_path
        subject, exam_type = _infer_subject_and_type(
            command_path or "",
            config.get("subject"),
            config.get("exam_type"),
        )
        visible_at = _parse_dt(config.get("visible_at"))
        starts_at = _parse_dt(config.get("starts_at"))
        if visible_at and starts_at and starts_at < visible_at:
            raise ValueError("starts_at visible_at dan oldin bo'lishi mumkin emas")
        base_price = max(0, _as_int(config.get("base_price"), 0))
        discount = _as_int(config.get("share_discount_percent"), 0)
        discount = max(0, min(100, discount))
        duration = _as_int(config.get("duration_minutes"), 20)
        if duration <= 0:
            duration = 20

        questions = []
        errors = []
        for offset, row in enumerate(_sheet_records(question_sheet), start=2):
            try:
                text = _as_str(row.get("question_text"))
                if not text:
                    raise ValueError("question_text bo'sh")
                options = [
                    _as_str(row.get("opt_a") or row.get("option_a")),
                    _as_str(row.get("opt_b") or row.get("option_b")),
                    _as_str(row.get("opt_c") or row.get("option_c")),
                    _as_str(row.get("opt_d") or row.get("option_d")),
                ]
                if any(opt is None for opt in options):
                    raise ValueError("opt_a, opt_b, opt_c, opt_d to'ldirilishi kerak")
                order_num = _as_int(row.get("order_num"), offset - 1)
                questions.append({
                    "order_num": order_num,
                    "question_text": text,
                    "image_top": _as_str(row.get("image_top")),
                    "image_bottom": _as_str(row.get("image_bottom")),
                    "option_a": options[0],
                    "option_b": options[1],
                    "option_c": options[2],
                    "option_d": options[3],
                    "correct_option": _correct_index(row.get("correct") if "correct" in row else row.get("correct_option")),
                    "difficulty_beta": _clamp_beta(row.get("beta") if "beta" in row else row.get("difficulty_beta")),
                    "option_sources": _option_sources(row),
                    "solution_text": _as_str(row.get("solution_text")),
                    "video_solution_url": _as_str(row.get("video_url") or row.get("video_solution_url")),
                })
            except ValueError as exc:
                errors.append(f"{offset}-qator: {exc}")
        if errors:
            preview = "\n".join(errors[:12])
            extra = f"\n... va yana {len(errors) - 12} ta" if len(errors) > 12 else ""
            raise ValueError(preview + extra)
        if not questions:
            raise ValueError("QUESTIONS sahifasida savol yo'q")
        questions.sort(key=lambda item: (item["order_num"], ))
        return {
            "command_path": command_path,
            "title": title,
            "subject": subject,
            "exam_type": exam_type,
            "visible_at": visible_at,
            "starts_at": starts_at,
            "base_price": base_price,
            "free_for_courses": _clean_courses(config.get("free_for_courses")),
            "share_discount_percent": discount,
            "full_analysis_video": _as_str(config.get("full_analysis_video")),
            "duration_minutes": duration,
            "link": _as_str(config.get("link")),
            "analysis_text": _as_str(config.get("analysis_text")),
            "questions": questions,
        }
    finally:
        wb.close()


async def save_quiz_workbook(path: str) -> dict:
    payload = parse_quiz_workbook(path)
    async with async_session() as session:
        quiz = None
        updated = False
        if payload["command_path"]:
            quiz = (
                await session.execute(
                    select(ExamQuiz)
                    .where(ExamQuiz.command_path == payload["command_path"])
                    .order_by(ExamQuiz.id.desc())
                )
            ).scalars().first()
        if quiz is None:
            quiz = ExamQuiz(
                title=payload["title"],
                subject=payload["subject"],
                exam_type=payload["exam_type"],
                is_active=True,
            )
            session.add(quiz)
            await session.flush()
        else:
            updated = True
            await session.execute(delete(ExamQuestion).where(ExamQuestion.quiz_id == quiz.id))

        quiz.title = payload["title"]
        quiz.subject = payload["subject"]
        quiz.exam_type = payload["exam_type"]
        quiz.command_path = payload["command_path"]
        quiz.visible_at = payload["visible_at"]
        quiz.starts_at = payload["starts_at"]
        quiz.base_price = payload["base_price"]
        quiz.price = payload["base_price"]
        quiz.free_for_courses = payload["free_for_courses"]
        quiz.share_discount_percent = payload["share_discount_percent"]
        quiz.full_analysis_video = payload["full_analysis_video"]
        quiz.duration_minutes = payload["duration_minutes"]
        quiz.link = payload["link"]
        quiz.analysis_text = payload["analysis_text"]
        quiz.announced_at = payload["visible_at"] or datetime.now()
        quiz.is_active = True
        await session.flush()
        quiz.pdf_url = f"/api/exams/paper/{quiz.id}?kind=test"
        quiz.answer_sheet_url = f"/api/exams/paper/{quiz.id}?kind=answers"

        for item in payload["questions"]:
            sources = item["option_sources"]
            session.add(ExamQuestion(
                quiz_id=quiz.id,
                order_num=item["order_num"],
                question_text=item["question_text"],
                image_top=item["image_top"],
                image_bottom=item["image_bottom"],
                option_a=item["option_a"],
                option_b=item["option_b"],
                option_c=item["option_c"],
                option_d=item["option_d"],
                correct_option=item["correct_option"],
                difficulty_beta=item["difficulty_beta"],
                option_sources=sources,
                book_links_json=sources,
                solution_text=item["solution_text"],
                analysis_text=item["solution_text"],
                video_solution_url=item["video_solution_url"],
                difficulty="orta",
            ))
        await session.commit()
        return {
            "quiz_id": quiz.id,
            "title": quiz.title,
            "questions": len(payload["questions"]),
            "updated": updated,
            "command_path": quiz.command_path,
            "subject": quiz.subject,
            "exam_type": quiz.exam_type,
            "base_price": quiz.base_price or 0,
        }


def _summary_text(summary: dict) -> str:
    action = "yangilandi" if summary["updated"] else "saqlandi"
    path = html.escape(summary["command_path"] or "—")
    title = html.escape(summary["title"])
    return (
        f"✅ Test {action}.\n"
        f"<b>{title}</b>\n"
        f"ID: {summary['quiz_id']}\n"
        f"Bo'lim: <code>{path}</code>\n"
        f"Fan: {html.escape(summary['subject'])} · tur: {html.escape(summary['exam_type'])}\n"
        f"Savollar: {summary['questions']}\n"
        f"Narx: {summary['base_price']} so'm\n\n"
        "Vaqtlar server soati bo'yicha solishtiriladi."
    )


async def _download_document(message: Message, bot: Bot) -> str:
    os.makedirs("downloads", exist_ok=True)
    filename = message.document.file_name or "quiz.xlsx"
    safe_name = os.path.basename(filename)
    file_path = os.path.join("downloads", safe_name)
    file = await bot.get_file(message.document.file_id)
    await bot.download_file(file.file_path, file_path)
    return file_path


def _is_excel_name(name: str) -> bool:
    lower = (name or "").lower()
    return lower.endswith(".xlsx") or lower.endswith(".xlsm")


@admin_quiz_router.message(Command("import_quiz"), F.from_user.id.in_(QUIZ_ADMINS))
async def cmd_import_quiz(message: Message, state: FSMContext):
    await state.set_state(ImportQuizState.waiting_file)
    await message.answer(
        "📥 Namunaviy test yoki milliy sertifikat Excelini yuboring.\n\n"
        "Sahifalar: <b>CONFIG</b> va <b>QUESTIONS</b>.\n"
        "CONFIG: command_path, quiz_title, visible_at, starts_at, base_price, "
        "free_for_courses, share_discount_percent, full_analysis_video.\n"
        "Ixtiyoriy: subject, exam_type, duration_minutes.\n"
        "QUESTIONS: order_num, question_text, image_top, image_bottom, "
        "opt_a, opt_b, opt_c, opt_d, correct, beta, ref_a, ref_b, ref_c, ref_d, "
        "solution_text, video_url.",
        parse_mode="HTML",
    )


@admin_quiz_router.message(ImportQuizState.waiting_file, F.document, F.from_user.id.in_(QUIZ_ADMINS))
async def import_quiz_waiting_file(message: Message, bot: Bot, state: FSMContext):
    if not _is_excel_name(message.document.file_name or ""):
        await message.answer("❌ .xlsx fayl yuboring. Bekor qilish uchun /import_quiz ni qayta yuboring.")
        return
    status = await message.answer("📥 Fayl o'qilmoqda...")
    file_path = await _download_document(message, bot)
    try:
        if not is_quiz_workbook(file_path):
            await status.edit_text("❌ Faylda CONFIG va QUESTIONS sahifalari topilmadi.")
            return
        summary = await save_quiz_workbook(file_path)
        await state.clear()
        await status.edit_text(_summary_text(summary), parse_mode="HTML")
    except Exception as exc:
        await status.edit_text(f"❌ Import xatosi:\n{html.escape(str(exc))}", parse_mode="HTML")
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)


@admin_quiz_router.message(StateFilter(None), F.document, F.from_user.id.in_(QUIZ_ADMINS))
async def import_quiz_excel(message: Message, bot: Bot):
    if not _is_excel_name(message.document.file_name or ""):
        raise SkipHandler()
    file_path = await _download_document(message, bot)
    try:
        if not is_quiz_workbook(file_path):
            raise SkipHandler()
        status = await message.answer("📥 Test import qilinmoqda...")
        summary = await save_quiz_workbook(file_path)
        await status.edit_text(_summary_text(summary), parse_mode="HTML")
    except SkipHandler:
        raise
    except Exception as exc:
        await message.answer(f"❌ Import xatosi:\n{html.escape(str(exc))}", parse_mode="HTML")
    finally:
        if file_path and os.path.exists(file_path):
            os.remove(file_path)


@admin_quiz_router.message(F.text.regexp(r"^/[\w./-]+"))
async def open_quiz_section(message: Message, state: FSMContext):
    current = await state.get_state()
    if current:
        raise SkipHandler()
    raw = (message.text or "").split()[0].split("@")[0].strip()
    if raw in {"/start", "/help", "/import_quiz", "/add_news", "/admin"}:
        raise SkipHandler()
    now = datetime.now()
    async with async_session() as session:
        rows = (
            await session.execute(
                select(ExamQuiz).where(
                    ExamQuiz.is_active == True,
                    ExamQuiz.command_path == raw,
                ).order_by(ExamQuiz.id.desc())
            )
        ).scalars().all()
    visible = [quiz for quiz in rows if not quiz.visible_at or quiz.visible_at <= now]
    if not visible:
        raise SkipHandler()

    lines = [f"📚 <b>{html.escape(raw)}</b> bo'limidagi testlar:\n"]
    for quiz in visible[:8]:
        if quiz.starts_at and now < quiz.starts_at:
            delta = quiz.starts_at - now
            hours = int(delta.total_seconds() // 3600)
            days, hour_rem = divmod(hours, 24)
            when = f"{days} kun {hour_rem} soatdan keyin" if days else f"{hour_rem} soatdan keyin"
            state_label = f"e'lon qilingan, boshlanishiga {when}"
        else:
            state_label = "ochiq"
        price = int(quiz.base_price or 0) or int(quiz.price or 0)
        price_label = "bepul" if price <= 0 else f"{price} so'm"
        lines.append(f"• <b>{html.escape(quiz.title)}</b>\n  {price_label} · {state_label}")
    keyboard = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="Testlar bo'limini ochish", web_app=WebAppInfo(url=WEBAPP_URL))
    ]])
    await message.answer("\n".join(lines), parse_mode="HTML", reply_markup=keyboard)
